"""End-to-end demo: one content document, two house styles.

Illustrates the process the library implements today:

  Markdown (the editable surface)
      -> block tree (the canonical form)          [from_markdown]
      -> the same tree, rendered against a base       [render]
         document's own named styles                  -> .docx

The point of the two houses is the design's core bet (ADR-0001): the content
document knows nothing about appearance, so the *same* tree renders correctly
into completely different style vocabularies. Structure is decoupled from look.

Run it from the repository root:

    python demo/run_demo.py

Generated files land in demo/out/ (gitignored). No template or client data is
involved; the base documents are built here with plain, invented style names.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

# Make the package importable without an install.
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from docx import Document  # noqa: E402
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402

import template_engine as te  # noqa: E402

OUT = ROOT / "demo" / "out"

# Two houses. Same roles, different style names, as two clients' templates would
# differ. (One uses German display names, one terse client codes.)
HAEUSER = {
    "haus-de": {
        "ueberschrift_1": "Titel 1",
        "ueberschrift_2": "Titel 2",
        "fliesstext": "Fliesstext",
        "aufzaehlung": "Liste Punkte",
        "nummerierte_liste": "Liste Nummer",
    },
    "haus-code": {
        "ueberschrift_1": "H1",
        "ueberschrift_2": "H2",
        "fliesstext": "Body",
        "aufzaehlung": "Bullet",
        "nummerierte_liste": "Number",
    },
}


def basis_docx(rollen: dict[str, str], kopf: str) -> bytes:
    """A stand-in base document carrying the house's named styles and a header.

    In production this is the client's own template; here it is synthesised so
    the demo runs anywhere.
    """
    doc = Document()
    for stilname in dict.fromkeys(rollen.values()):
        doc.styles.add_style(stilname, WD_STYLE_TYPE.PARAGRAPH)
    doc.sections[0].header.paragraphs[0].text = kopf
    doc.add_paragraph("Platzhalter, der beim Rendern verschwindet.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def zeige_baum(dok: te.Dokument) -> None:
    print("  metadata:", dok.metadaten)
    for b in dok.bloecke:
        if isinstance(b, te.Ueberschrift):
            print(f"    Ueberschrift(ebene={b.ebene}) {b.text!r}")
        elif isinstance(b, te.Absatz):
            print(f"    Absatz {b.text[:50]!r}")
        elif isinstance(b, te.Aufzaehlung):
            print(f"    Aufzaehlung {b.punkte}")
        elif isinstance(b, te.NummerierteListe):
            print(f"    NummerierteListe {b.punkte}")


def paragraphen(docx: bytes) -> list[tuple[str, str]]:
    doc = Document(io.BytesIO(docx))
    return [(p.text, p.style.name) for p in doc.paragraphs]


def main() -> None:
    OUT.mkdir(exist_ok=True)
    quelle = (ROOT / "demo" / "content.md").read_text(encoding="utf-8")

    print("STEP 1  Markdown -> block tree (the canonical form)")
    dok = te.from_markdown(quelle)
    zeige_baum(dok)

    print("\nSTEP 2  round-trip guarantee: tree -> Markdown -> tree is identity")
    dok2 = te.from_markdown(te.to_markdown(dok))
    print("  stable:", dok2 == dok)

    print("\nSTEP 3+4  render the SAME tree into each house's own styles")
    for haus, rollen in HAEUSER.items():
        basis = basis_docx(rollen, kopf=f"{haus.upper()} BRIEFKOPF")
        docx = te.render(dok, te.StyleProfile(rollen), basis)
        ziel = OUT / f"angebot__{haus}.docx"
        ziel.write_bytes(docx)
        print(f"\n  [{haus}] -> {ziel.relative_to(ROOT)}")
        for text, stil in paragraphen(docx):
            print(f"      {stil:14} | {text}")

    print("\nSame content, two vocabularies: structure is decoupled from look.")


if __name__ == "__main__":
    main()
