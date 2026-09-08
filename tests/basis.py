"""Synthetic base documents for renderer tests.

Real client templates are not needed to test the renderer: what matters is that
it follows *the base document's own* style names, whatever they happen to be.
So these fixtures deliberately vary the style vocabulary, from tidy German
display names to underscore-style client codes, including a non-ASCII name. A
renderer that passes against all of them is not quietly assuming a fixed naming
scheme, which is the whole point of ADR-0001 and of the language-dependent
style-name open question in the SPEC.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

# The roles this catalogue uses. Every base must provide a style for each.
ROLLEN = ("ueberschrift_1", "ueberschrift_2", "fliesstext", "aufzaehlung", "nummerierte_liste")


@dataclass(frozen=True)
class Basis:
    name: str
    rollen: dict[str, str]  # role -> style name, i.e. the profile mapping
    docx: bytes


def _baue(
    rollen: dict[str, str],
    header: str,
    footer: str | None,
    altinhalt: list[str],
) -> bytes:
    doc = Document()
    for stilname in dict.fromkeys(rollen.values()):  # unique, order-preserving
        doc.styles.add_style(stilname, WD_STYLE_TYPE.PARAGRAPH)
    doc.sections[0].header.paragraphs[0].text = header
    if footer is not None:
        doc.sections[0].footer.paragraphs[0].text = footer
    for zeile in altinhalt:
        doc.add_paragraph(zeile)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def alle_basen() -> list[Basis]:
    varianten = [
        (
            "german-display-names",
            {
                "ueberschrift_1": "Titel 1",
                "ueberschrift_2": "Titel 2",
                "fliesstext": "Fliesstext",
                "aufzaehlung": "Aufzählungsliste",
                "nummerierte_liste": "Nummerierte Liste",
            },
            "ACME LETTERHEAD",
            "Vertraulich",
            ["alter Absatz eins", "alter Absatz zwei"],
        ),
        (
            "client-code-names",
            {
                "ueberschrift_1": "CX_H1",
                "ueberschrift_2": "CX_H2",
                "fliesstext": "CX_Body",
                "aufzaehlung": "CX_Bullet",
                "nummerierte_liste": "CX_Num",
            },
            "Client X",
            None,
            ["prior content"],
        ),
        (
            "reused-style-for-both-lists",
            {
                "ueberschrift_1": "H1",
                "ueberschrift_2": "H2",
                "fliesstext": "P",
                # A base may use one list style for both list kinds.
                "aufzaehlung": "Liste",
                "nummerierte_liste": "Liste",
            },
            "Minimal",
            "Seite",
            [],
        ),
    ]
    return [Basis(name, rollen, _baue(rollen, h, f, alt)) for name, rollen, h, f, alt in varianten]
