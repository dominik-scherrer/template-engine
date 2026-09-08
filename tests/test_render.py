"""Renderer tests: content rendered against a base document's named styles.

The base document is synthesised here with python-docx, so these tests need no
real client template. It carries a few custom paragraph styles, a header, and
some pre-existing body content, which is exactly what the renderer must clear
while keeping the styles and the header.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

import template_engine as te
from template_engine import FehlenderStil, StyleProfile, UnbekannteRolle, render


def _basis_docx(
    stile: list[str],
    header: str = "ACME LETTERHEAD",
    altinhalt: str = "vorbestehender Rumpfinhalt",
) -> bytes:
    doc = Document()
    for name in stile:
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    doc.sections[0].header.paragraphs[0].text = header
    doc.add_paragraph(altinhalt)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _paragraphen(docx: bytes) -> list[tuple[str, str]]:
    doc = Document(io.BytesIO(docx))
    return [(p.text, p.style.name) for p in doc.paragraphs]


def test_rendert_gegen_benannte_stile() -> None:
    basis = _basis_docx(["Titel 1", "Titel 2", "Fliesstext"])
    profil = StyleProfile(
        {
            "ueberschrift_1": "Titel 1",
            "ueberschrift_2": "Titel 2",
            "fliesstext": "Fliesstext",
        }
    )
    dok = te.Dokument(
        bloecke=[
            te.Ueberschrift(1, "Ausgangslage"),
            te.Absatz("Ein erster Absatz."),
            te.Ueberschrift(2, "Vorgehen"),
            te.Absatz("Ein zweiter Absatz."),
        ]
    )
    out = render(dok, profil, basis)

    assert _paragraphen(out) == [
        ("Ausgangslage", "Titel 1"),
        ("Ein erster Absatz.", "Fliesstext"),
        ("Vorgehen", "Titel 2"),
        ("Ein zweiter Absatz.", "Fliesstext"),
    ]


def test_kopfzeile_ueberlebt_und_alter_inhalt_verschwindet() -> None:
    basis = _basis_docx(["Titel 1", "Fliesstext"], header="ACME LETTERHEAD")
    profil = StyleProfile({"ueberschrift_1": "Titel 1", "fliesstext": "Fliesstext"})
    dok = te.Dokument(bloecke=[te.Ueberschrift(1, "Neu")])
    out = render(dok, profil, basis)

    doc = Document(io.BytesIO(out))
    assert doc.sections[0].header.paragraphs[0].text == "ACME LETTERHEAD"
    assert all(p.text != "vorbestehender Rumpfinhalt" for p in doc.paragraphs)


def test_unbekannte_rolle_wird_laut() -> None:
    basis = _basis_docx(["Titel 1", "Fliesstext"])
    profil = StyleProfile({"ueberschrift_1": "Titel 1", "fliesstext": "Fliesstext"})
    dok = te.Dokument(bloecke=[te.Ueberschrift(3, "Zu tief")])  # ueberschrift_3 unmapped
    with pytest.raises(UnbekannteRolle):
        render(dok, profil, basis)


def test_fehlender_stil_wird_laut() -> None:
    basis = _basis_docx(["Fliesstext"])  # no "Titel 1" in the base document
    profil = StyleProfile({"ueberschrift_1": "Titel 1", "fliesstext": "Fliesstext"})
    dok = te.Dokument(bloecke=[te.Ueberschrift(1, "X"), te.Absatz("y")])
    with pytest.raises(FehlenderStil):
        render(dok, profil, basis)


def test_pruefung_vor_mutation() -> None:
    # A render that fails conformance must raise rather than return a document.
    basis = _basis_docx(["Fliesstext"])
    profil = StyleProfile({"ueberschrift_1": "Fehlt", "fliesstext": "Fliesstext"})
    dok = te.Dokument(bloecke=[te.Absatz("ok"), te.Ueberschrift(1, "kaputt")])
    with pytest.raises(FehlenderStil):
        render(dok, profil, basis)


def test_rolle_fuer_block() -> None:
    assert te.rolle_fuer_block(te.Ueberschrift(2, "x")) == "ueberschrift_2"
    assert te.rolle_fuer_block(te.Absatz("x")) == "fliesstext"


def test_profil_json_rundreise() -> None:
    p = StyleProfile({"ueberschrift_1": "Titel 1", "fliesstext": "Fliesstext"})
    assert StyleProfile.from_json(p.to_json()) == p
