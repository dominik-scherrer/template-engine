"""Ingest tests: deterministic reading and the proposer seam.

The LLM proposer is tested with its HTTP call substituted, so no network or key
is needed; the live call is the caller's to exercise.
"""

from __future__ import annotations

import io
import zipfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from template_engine import ingest


def _synth(styles: list[str], *, dotx: bool) -> bytes:
    doc = Document()
    for n in styles:
        doc.styles.add_style(n, WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("Muster AG", style=styles[0] if styles else "Normal")
    doc.add_paragraph("Ein Absatz.", style="Normal")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    if not dotx:
        return data
    zin = zipfile.ZipFile(io.BytesIO(data))
    ct = zin.read("[Content_Types].xml").decode().replace(ingest.CT_DOCUMENT, ingest.CT_TEMPLATE)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for it in zin.infolist():
            z.writestr(
                it, ct.encode() if it.filename == "[Content_Types].xml" else zin.read(it.filename)
            )
    return out.getvalue()


def test_dotx_wird_oeffenbar() -> None:
    dotx = _synth(["Titel 1"], dotx=True)
    # python-docx rejects the raw .dotx, accepts the patched bytes
    Document(io.BytesIO(ingest.als_dokument_bytes(dotx)))


def test_stil_inventar_und_absatzfolge() -> None:
    data = _synth(["Titel 1", "Fliesstext"], dotx=False)
    inv = ingest.stil_inventar(data)
    ids = {s["styleId"] for s in inv}
    assert "Titel1" in ids and "Fliesstext" in ids
    folge = ingest.absatzfolge(data)
    assert folge[0]["text"] == "Muster AG"
    assert folge[0]["styleId"] == "Titel1"


def test_heuristischer_vorschlag() -> None:
    kontext = ingest.kontext_aus_bytes(_synth(["Titel 1", "Fliesstext"], dotx=False))
    v = ingest.heuristischer_vorschlag(kontext)
    assert v.rollen["ueberschrift_1"] in {"Titel1", "Heading1"}  # a heading style
    assert v.rollen["fliesstext"] == "Fliesstext"


def test_llm_konfig_from_env(monkeypatch) -> None:
    monkeypatch.delenv("TEMPLATE_ENGINE_LLM_BASE_URL", raising=False)
    assert ingest.LLMKonfig.from_env() is None
    monkeypatch.setenv("TEMPLATE_ENGINE_LLM_BASE_URL", "https://x/v1")
    monkeypatch.setenv("TEMPLATE_ENGINE_LLM_API_KEY", "k")
    monkeypatch.setenv("TEMPLATE_ENGINE_LLM_MODEL", "m")
    k = ingest.LLMKonfig.from_env()
    assert k and k.base_url == "https://x/v1" and k.model == "m"


def test_llm_proposer_parst_antwort(monkeypatch) -> None:
    canned = (
        'Here you go: {"rollen": {"fliesstext": "Fliesstext"}, '
        '"bindestellen": [{"absatz": 0, "feld": "kunde_name", "text": "Muster AG"}], '
        '"koerpergrenze": 5}'
    )
    monkeypatch.setattr(ingest, "_chat", lambda konfig, messages: canned)
    proposer = ingest.LLMProposer(ingest.LLMKonfig("https://x/v1", "k", "m"))
    kontext = ingest.kontext_aus_bytes(_synth(["Fliesstext"], dotx=False))
    v = proposer.vorschlag(kontext)
    assert v.rollen == {"fliesstext": "Fliesstext"}
    assert v.bindestellen[0]["feld"] == "kunde_name"
    assert v.koerpergrenze == 5


def test_llm_proposer_meldet_fehler(monkeypatch) -> None:
    def boom(konfig, messages):
        raise ConnectionError("no route")

    monkeypatch.setattr(ingest, "_chat", boom)
    proposer = ingest.LLMProposer(ingest.LLMKonfig("https://x/v1", "k", "m"))
    kontext = ingest.kontext_aus_bytes(_synth(["Fliesstext"], dotx=False))
    try:
        proposer.vorschlag(kontext)
    except ingest.LLMFehler:
        return
    raise AssertionError("expected LLMFehler")
